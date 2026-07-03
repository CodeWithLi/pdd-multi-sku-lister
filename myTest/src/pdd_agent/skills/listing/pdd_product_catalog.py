"""拼多多所有商品.xls — 商品编码 / 颜色规格 / 商品类型。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import xlrd

def default_catalog_path() -> Path:
    import pdd_agent.paths as paths

    return paths.PROJECT_ROOT / "data" / "assets" / "拼多多所有商品.xls"


@dataclass(frozen=True)
class PddProductEntry:
    code: str
    color_spec: str
    product_type: str
    inventory: int = 0
    price: float = 0.0


# 规则无法给出有意义分组时的「弱类型」（如亚克力配件需再推断子类）
WEAK_SUB_CATEGORIES = frozenset({"", "亚克力", "半成品", "配件"})

# 预览图文件名含 3 段及以上「+」时，型号统一解析为此编码（展示名为 {母商品类型}+套装）
BUNDLE_MODEL_CODE = "套装"

# 颜色规格关键词 → 尺寸子类（按优先级，先匹配先生效）
_SUB_CATEGORY_KEYWORDS: tuple[tuple[str, str], ...] = (
    ("挂件", "挂件"),
    ("套尺", "套尺"),
    ("磁吸夹", "磁吸夹"),
    ("冰箱贴", "冰箱贴"),
    ("PP夹", "PP夹"),
    ("发夹", "发夹"),
    ("钥匙扣", "钥匙扣"),
    ("打卡棒", "打卡棒"),
    ("书签", "书签"),
    ("便签本", "便签"),
    ("便签", "便签"),
    ("徽章", "徽章"),
    ("摇摇乐", "尺子"),
    ("直尺", "尺子"),
    ("尺子", "尺子"),
    ("卡套", "卡套"),
    ("贴纸", "贴纸"),
    ("镜子", "镜子"),
    ("镜", "镜子"),
)


def _norm_code(value: str) -> str:
    return str(value or "").strip().upper()


_ALPHANUM_PREFIX = re.compile(r"^([A-Z0-9]+)")


def product_code_alphanumeric_prefix(code: str) -> str:
    """商品编码 leading 段（不含中文后缀），如 BD037A奶油泡芙笔袋 → BD037A。"""
    match = _ALPHANUM_PREFIX.match(_norm_code(code))
    return match.group(1) if match else _norm_code(code)


def product_code_family_key(code: str) -> str:
    """去掉末位变体字母后的商品族键：BD037A→BD037，YKL001N→YKL001。"""
    prefix = product_code_alphanumeric_prefix(code)
    if len(prefix) >= 2 and prefix[-1].isalpha():
        return prefix[:-1]
    return prefix


def find_catalog_entry(
    catalog: dict[str, PddProductEntry],
    code: str,
) -> PddProductEntry | None:
    """按飞书商品编码查找（保留英文+中文全名，支持前缀互认）。"""
    from pdd_agent.integrations.feishu.listing_feishu_catalog import catalog_entry

    return catalog_entry(catalog, code)


def catalog_codes_match(left: str, right: str) -> bool:
    """两侧编码是否同一商品（全名 / 前缀 / 归一化，不截断中文）。"""
    from pdd_agent.integrations.feishu.listing_feishu_catalog import codes_match

    return codes_match(left, right)


def _parse_cell_int(value, default: int = 0) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def _parse_cell_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def _format_price(value: float) -> str:
    rounded = round(value, 2)
    if rounded == int(rounded):
        return str(int(rounded))
    return f"{rounded:.2f}".rstrip("0").rstrip(".")


def _load_xls_sheet(excel_path: Path):
    book = xlrd.open_workbook(str(excel_path))
    if book.nsheets <= 0:
        return None, []
    sheet = book.sheet_by_index(0)
    if sheet.nrows <= 1:
        return None, []
    headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
    rows: list[tuple] = []
    for row_idx in range(1, sheet.nrows):
        rows.append(tuple(sheet.cell_value(row_idx, col) for col in range(sheet.ncols)))
    return headers, rows


def _load_xlsx_sheet(excel_path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active
    iter_rows = ws.iter_rows(values_only=True)
    header_row = next(iter_rows, None)
    if not header_row:
        wb.close()
        return None, []
    headers = [str(cell or "").strip() for cell in header_row]
    rows = [row for row in iter_rows if row and row[0]]
    wb.close()
    return headers, rows


def _load_sheet_rows(excel_path: Path) -> tuple[list[str], list[tuple]] | tuple[None, list]:
    suffix = excel_path.suffix.lower()
    if suffix == ".xls":
        return _load_xls_sheet(excel_path)
    if suffix == ".xlsx":
        return _load_xlsx_sheet(excel_path)
    return None, []


def load_pdd_product_catalog(path: Path | None = None) -> dict[str, PddProductEntry]:
    excel_path = path or default_catalog_path()
    if not excel_path.is_file():
        return {}

    try:
        headers, rows = _load_sheet_rows(excel_path)
    except Exception:
        return {}
    if not headers:
        return {}

    code_col = headers.index("商品编码") if "商品编码" in headers else 0
    color_col = headers.index("颜色规格") if "颜色规格" in headers else 1
    type_col = headers.index("商品类型") if "商品类型" in headers else 2
    inv_col = headers.index("库存") if "库存" in headers else None
    price_col = None
    for col_name in ("拼多多定价+5元", "拼多多定价", "价格"):
        if col_name in headers:
            price_col = headers.index(col_name)
            break

    catalog: dict[str, PddProductEntry] = {}
    for row in rows:
        if not row or code_col >= len(row):
            continue
        raw_code = str(row[code_col] or "").strip()
        code = _norm_code(raw_code)
        if not code:
            continue
        color_spec = str(row[color_col] or "").strip() if color_col < len(row) else ""
        if not color_spec:
            continue
        catalog[code] = PddProductEntry(
            code=code,
            color_spec=color_spec,
            product_type=str(row[type_col] or "").strip() if type_col < len(row) else "",
            inventory=_parse_cell_int(row[inv_col], 0) if inv_col is not None and inv_col < len(row) else 0,
            price=_parse_cell_float(row[price_col], 0.0) if price_col is not None and price_col < len(row) else 0.0,
        )
    return catalog


@lru_cache
def get_pdd_product_catalog() -> dict[str, PddProductEntry]:
    from pdd_agent.integrations.feishu.pdd_product_catalog_sync import load_unified_pdd_product_catalog

    return load_unified_pdd_product_catalog()


def resolve_listing_price_scope(*, family_code: str = "", image_set: str = "", listing_label: str = "") -> str:
    family = str(family_code or "").strip()
    image = str(image_set or "").strip()
    if family and image:
        return f"{family}/{image}"
    return str(listing_label or "").strip()


class ComboPricingError(Exception):
    """商品表缺编码或定价无效，无法计算拼单价/单买价。"""

    def __init__(
        self,
        *,
        missing_codes: list[str] | None = None,
        invalid_codes: list[str] | None = None,
        reason: str = "catalog_code_missing",
    ) -> None:
        self.missing_codes = [str(code).strip() for code in (missing_codes or []) if str(code).strip()]
        self.invalid_codes = [str(code).strip() for code in (invalid_codes or []) if str(code).strip()]
        self.reason = str(reason or "catalog_code_missing").strip() or "catalog_code_missing"
        codes = self.invalid_codes or self.missing_codes
        label = "、".join(codes)
        if self.reason == "pricing_invalid":
            super().__init__(f"商品表编码[{label}]拼多多定价无效")
        elif self.reason == "jushuitan_cost_missing":
            super().__init__(f"聚水潭成本缺失[{label}]")
        else:
            super().__init__(f"商品表无编码[{label}]，无法定价")


def resolve_combo_pricing(
    style_code: str,
    model_code: str,
    catalog: dict[str, PddProductEntry],
    defaults: dict[str, str],
    *,
    listing_scope: str = "",
    bundle_components: list[str] | None = None,
) -> tuple[str, str, str]:
    """组合 SKU 库存/拼单价/单买价（定价来自商品表；缺编码或定价无效时报错）。"""
    from pdd_agent.skills.listing.listing_price_float import get_listing_price_float_registry

    _ = defaults  # 保留入参兼容；不再使用 199.9/299.9 兜底
    fallback_inv = "10"
    registry = get_listing_price_float_registry()
    scope = str(listing_scope or "").strip()

    def _require_entry(code: str) -> PddProductEntry:
        entry = find_catalog_entry(catalog, code)
        if not entry:
            raise ComboPricingError(missing_codes=[code], reason="catalog_code_missing")
        if float(entry.price or 0.0) <= 0:
            raise ComboPricingError(invalid_codes=[code], reason="pricing_invalid")
        return entry

    def _component_price(entry: PddProductEntry) -> float:
        base = float(entry.price or 0.0)
        if scope and registry.enabled:
            return registry.effective_price(scope, base)
        return base

    def _finalize(style_entry: PddProductEntry, group_price: float, inventory: int | str) -> tuple[str, str, str]:
        gp = round(float(group_price), 2)
        if gp <= 0:
            raise ComboPricingError(invalid_codes=[style_entry.code], reason="pricing_invalid")
        sp = registry.effective_single_price(scope, gp)
        inv = str(inventory if inventory is not None else style_entry.inventory or fallback_inv)
        return inv, _format_price(gp), _format_price(sp)

    if catalog_codes_match(model_code, style_code):
        style_entry = _require_entry(style_code)
        group_price = _component_price(style_entry)
        return _finalize(style_entry, group_price, style_entry.inventory)

    if _norm_code(model_code) == BUNDLE_MODEL_CODE and bundle_components:
        component_codes: list[str] = []
        for code in bundle_components:
            entry = find_catalog_entry(catalog, code)
            if entry and entry.code not in component_codes:
                component_codes.append(entry.code)
        if not component_codes:
            raise ComboPricingError(missing_codes=[style_code], reason="catalog_code_missing")
        entries = [_require_entry(code) for code in component_codes]
        inventories = [int(entry.inventory or 0) for entry in entries]
        inventory = min(inventories) if inventories else fallback_inv
        group_price = sum(_component_price(entry) for entry in entries)
        return _finalize(entries[0], group_price, inventory)

    style_entry = _require_entry(style_code)
    model_entry = _require_entry(model_code)
    inventory = min(int(style_entry.inventory or 0), int(model_entry.inventory or 0))
    group_price = _component_price(style_entry) + _component_price(model_entry)
    return _finalize(style_entry, group_price, inventory)


def catalog_has_code(code: str) -> bool:
    return find_catalog_entry(get_pdd_product_catalog(), code) is not None


def get_catalog_entry(code: str) -> PddProductEntry | None:
    return find_catalog_entry(get_pdd_product_catalog(), code)


def lookup_catalog_spec(code: str, catalog: dict[str, PddProductEntry] | None = None) -> str | None:
    """颜色规格（等同原 sku+名称.xlsx 中的名称列）。"""
    entry = find_catalog_entry(catalog or get_pdd_product_catalog(), code)
    return entry.color_spec if entry else None


def is_weak_sub_category(category: str) -> bool:
    cat = str(category or "").strip()
    if not cat or cat in WEAK_SUB_CATEGORIES:
        return True
    return len(cat) > 8


def infer_sub_product_category_rules(entry: PddProductEntry) -> str | None:
    """仅用规则从商品表推断子类；无法推断时返回 None。"""
    spec = entry.color_spec
    for keyword, category in _SUB_CATEGORY_KEYWORDS:
        if keyword in spec:
            return category
    product_type = entry.product_type.strip()
    if product_type and product_type not in WEAK_SUB_CATEGORIES:
        return product_type
    return None


def fallback_sub_category_from_spec(entry: PddProductEntry) -> str:
    """无 AI 时的最后兜底：从颜色规格截取可读短名。"""
    spec = entry.color_spec.strip()
    for prefix in ("亚克力定制款", "亚克力定制", "定制款", "亚克力"):
        if spec.startswith(prefix):
            spec = spec[len(prefix) :].strip()
    spec = re.sub(r"[()（）【】].*", "", spec).strip("款 ")
    if spec:
        return spec[:6]
    return entry.code


def infer_sub_product_category(entry: PddProductEntry) -> str:
    """子商品归类名（仅规则，不含 AI）。"""
    ruled = infer_sub_product_category_rules(entry)
    if ruled:
        return ruled
    return fallback_sub_category_from_spec(entry)


def append_product_type_to_size_label(base_label: str, product_type: str) -> str:
    """在尺寸名后附加商品类型，如 笔袋+挂件 → 笔袋+挂件(公仔)。"""
    label = (base_label or "").strip()
    ptype = (product_type or "").strip()
    if not label or not ptype:
        return label
    if f"({ptype})" in label or label.endswith(f"({ptype})"):
        return label
    return f"{label}({ptype})"


def meaningful_product_type_for_label(entry: PddProductEntry | None) -> str:
    """可用于尺寸名后缀的商品类型（排除亚克力等弱类型）。"""
    if not entry:
        return ""
    ptype = entry.product_type.strip()
    if not ptype or ptype in WEAK_SUB_CATEGORIES:
        return ""
    return ptype


def resolve_size_label(
    style_code: str,
    model_code: str,
    catalog: dict[str, PddProductEntry],
    *,
    sub_category_cache: dict[str, str] | None = None,
) -> str:
    """第二个规格（尺寸）展示名：{母商品类型} / {母商品类型}+{子类} …"""
    base = find_catalog_entry(catalog, style_code)
    if not base:
        return model_code

    if catalog_codes_match(model_code, style_code):
        base_type = base.product_type.strip() or style_code
        return base_type
    if _norm_code(model_code) == BUNDLE_MODEL_CODE:
        base_type = base.product_type.strip() or style_code
        return f"{base_type}+{BUNDLE_MODEL_CODE}"

    sub = find_catalog_entry(catalog, model_code)
    if not sub:
        return model_code

    base_type = base.product_type.strip() or style_code
    cache = sub_category_cache or {}
    model_key = _norm_code(sub.code)
    sub_cat = cache.get(model_key) or cache.get(_norm_code(model_code))
    if not sub_cat:
        sub_cat = infer_sub_product_category(sub)
    return f"{base_type}+{sub_cat}"
