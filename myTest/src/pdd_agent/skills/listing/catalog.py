"""多规格素材目录解析 Skill。

当前主用结构：

data/assets/商品上架规格文件/
  BD046/                 # 母文件夹，一类商品
    第一套/
      BD046A/            # 实际上架编码
        1.png ... 5.png  # 主图
        6.png            # 占位，不上传
        7.png ... 9.png  # 详情图
        BD046A+AXBD001A.png  # SKU 预览图，文件名即组合编码

代码仍保留旧结构兼容，避免历史数据/测试立即失效。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from langchain_core.tools import StructuredTool
from openpyxl import load_workbook

from pdd_agent.core.config import YamlSettings
from pdd_agent.domain.models import ListingUnit, SkuCombo
from pdd_agent.skills.base import Skill

RESERVED_DIRS = frozenset({"主图", "详情页"})
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
PREVIEW_NAMES = ("预览图.png", "预览图.jpg", "预览图.jpeg", "预览图.webp")
IMAGE_SET_PATTERN = re.compile(r"^第.+套(图)?$")


@dataclass(frozen=True)
class SkuExcelInfo:
    code: str
    name: str
    inventory: int
    price: float


@dataclass(frozen=True)
class ListingTargetOption:
    selector: str
    family_code: str
    image_set: str
    product_codes: tuple[str, ...]


def _natural_sort_key(path: Path):
    parts = re.split(r"(\d+)", path.name.lower())
    return [int(p) if p.isdigit() else p for p in parts]


def _image_set_sort_key(name: str):
    cn_map = {
        "零": 0,
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
    }
    match = re.search(r"第([一二三四五六七八九十\d]+)套", name)
    if not match:
        return (1, name)
    token = match.group(1)
    if token.isdigit():
        return (0, int(token))
    if token in cn_map:
        return (0, cn_map[token])
    if token.startswith("十") and len(token) == 2 and token[1] in cn_map:
        return (0, 10 + cn_map[token[1]])
    return (0, sum(cn_map.get(ch, 0) for ch in token))


def is_image_set_dir(name: str) -> bool:
    return bool(IMAGE_SET_PATTERN.match(name.strip()))


def _parse_excel_int(value, default: int = 0) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def _parse_excel_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def _format_price(value: float) -> str:
    rounded = round(value, 2)
    if rounded == int(rounded):
        return str(int(rounded))
    return f"{rounded:.2f}".rstrip("0").rstrip(".")


def load_sku_excel_map(excel_path: Path) -> dict[str, SkuExcelInfo]:
    """读取 sku+名称.xlsx -> {SKU编码: SkuExcelInfo}。"""
    if not excel_path.is_file():
        return {}
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active
    mapping: dict[str, SkuExcelInfo] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if not name:
            continue
        mapping[code] = SkuExcelInfo(
            code=code,
            name=name,
            inventory=_parse_excel_int(row[2] if len(row) > 2 else None, 0),
            price=_parse_excel_float(row[3] if len(row) > 3 else None, 0.0),
        )
    wb.close()
    return mapping


def lookup_sku_name(code: str, excel_map: dict[str, SkuExcelInfo]) -> str | None:
    info = excel_map.get((code or "").strip())
    return info.name if info else None


def _load_excel_map_with_fallback(settings: YamlSettings, assets_root: Path) -> dict[str, SkuExcelInfo]:
    excel_map = load_sku_excel_map(settings.multi_spec.excel_path)
    if excel_map:
        return excel_map

    candidates = [
        assets_root / "sku+名称.xlsx",
        assets_root.parent / "sku+名称.xlsx",
        *sorted(assets_root.parent.glob("*/sku+*.xlsx")),
    ]
    seen: set[Path] = set()
    for candidate in candidates:
        candidate = candidate.resolve()
        if candidate in seen or candidate.name.startswith("~$"):
            continue
        seen.add(candidate)
        excel_map = load_sku_excel_map(candidate)
        if excel_map:
            return excel_map
    return {}


def resolve_combo_pricing(
    style_code: str,
    model_code: str,
    excel_map: dict[str, SkuExcelInfo],
    defaults: dict[str, str],
) -> tuple[str, str, str]:
    fallback_inv = defaults.get("inventory", "10")
    fallback_gp = defaults.get("group_price", "199.9")
    fallback_sp = defaults.get("single_price", "299.9")

    style = excel_map.get(style_code.strip())
    model = excel_map.get(model_code.strip())
    if not style or not model:
        return fallback_inv, fallback_gp, fallback_sp

    inventory = min(style.inventory, model.inventory)
    group_price = style.price + model.price
    single_price = group_price + 100
    return str(inventory), _format_price(group_price), _format_price(single_price)


def _collect_images(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    files = [f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in IMAGE_EXTS]
    return sorted(files, key=_natural_sort_key)


def _numbered_image_index(path: Path) -> int | None:
    if path.suffix.lower() not in IMAGE_EXTS:
        return None
    stem = path.stem.strip()
    return int(stem) if stem.isdigit() else None


def _collect_numbered_images(folder: Path, indexes: set[int]) -> list[Path]:
    if not folder.is_dir():
        return []
    files = [
        p
        for p in folder.iterdir()
        if p.is_file() and (idx := _numbered_image_index(p)) is not None and idx in indexes
    ]
    return sorted(files, key=lambda p: _numbered_image_index(p) or 0)


def _collect_flat_preview_images(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    files = [
        p
        for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in IMAGE_EXTS and _numbered_image_index(p) is None
    ]
    return sorted(files, key=_natural_sort_key)


def _preview_sort_key(path: Path, product_code: str):
    parsed = _parse_combo_code(path.stem, product_code)
    if parsed and parsed[0] == product_code and parsed[1] == product_code:
        return (0, _natural_sort_key(path))
    return (1, _natural_sort_key(path))


def _find_preview_image(combo_dir: Path) -> Path | None:
    for name in PREVIEW_NAMES:
        p = combo_dir / name
        if p.is_file():
            return p
    imgs = _collect_images(combo_dir)
    return imgs[0] if imgs else None


def _parse_combo_code(name: str, default_style_code: str = "") -> tuple[str, str] | None:
    stem = Path(name).stem.strip()
    if not stem:
        return None
    if "+" in stem:
        left, right = stem.split("+", 1)
        left, right = left.strip(), right.strip()
        if left and right:
            return left, right
        return None
    if default_style_code:
        return default_style_code, stem
    return stem, stem


def _new_layout_units(product_code: str, assets_root: Path) -> list[tuple[str, Path]]:
    """返回新结构中某个实际商品编码的 [(套图名, 编码目录)]。"""
    if not assets_root.is_dir():
        return []
    units: list[tuple[str, Path]] = []
    seen: set[Path] = set()

    for family_dir in sorted((p for p in assets_root.iterdir() if p.is_dir()), key=lambda p: p.name):
        for set_dir in sorted((p for p in family_dir.iterdir() if p.is_dir()), key=lambda p: _image_set_sort_key(p.name)):
            if not is_image_set_dir(set_dir.name):
                continue
            code_dir = set_dir / product_code
            if code_dir.is_dir() and code_dir not in seen:
                units.append((set_dir.name, code_dir))
                seen.add(code_dir)

    return units


def _new_layout_product_codes(assets_root: Path) -> list[str]:
    codes: set[str] = set()
    if not assets_root.is_dir():
        return []
    for family_dir in assets_root.iterdir():
        if not family_dir.is_dir():
            continue
        for set_dir in family_dir.iterdir():
            if not set_dir.is_dir() or not is_image_set_dir(set_dir.name):
                continue
            for code_dir in set_dir.iterdir():
                if code_dir.is_dir() and any(_numbered_image_index(p) is not None for p in code_dir.iterdir() if p.is_file()):
                    codes.add(code_dir.name)
    return sorted(codes)


def _code_dir_has_numbered_images(code_dir: Path) -> bool:
    return code_dir.is_dir() and any(
        _numbered_image_index(p) is not None for p in code_dir.iterdir() if p.is_file()
    )


def _new_layout_listing_targets(assets_root: Path) -> list[ListingTargetOption]:
    targets: list[ListingTargetOption] = []
    if not assets_root.is_dir():
        return targets
    for family_dir in sorted((p for p in assets_root.iterdir() if p.is_dir()), key=lambda p: p.name):
        for set_dir in sorted(
            (p for p in family_dir.iterdir() if p.is_dir() and is_image_set_dir(p.name)),
            key=lambda p: _image_set_sort_key(p.name),
        ):
            codes = tuple(
                sorted(
                    [
                        code_dir.name
                        for code_dir in set_dir.iterdir()
                        if _code_dir_has_numbered_images(code_dir)
                    ]
                )
            )
            if not codes:
                continue
            targets.append(
                ListingTargetOption(
                    selector=f"{family_dir.name}/{set_dir.name}",
                    family_code=family_dir.name,
                    image_set=set_dir.name,
                    product_codes=codes,
                )
            )
    return targets


def _parse_listing_target_selector(selector: str) -> tuple[str, str] | None:
    raw = str(selector or "").strip()
    if "/" not in raw:
        return None
    family_code, image_set = [part.strip() for part in raw.split("/", 1)]
    if not family_code or not image_set or not is_image_set_dir(image_set):
        return None
    return family_code, image_set


def list_image_sets(product_code: str, assets_root: Path) -> list[str]:
    new_units = _new_layout_units(product_code, assets_root)
    if new_units:
        return sorted({name for name, _ in new_units}, key=_image_set_sort_key)

    root = assets_root / product_code
    if not root.is_dir():
        return []
    return [
        d.name
        for d in sorted(root.iterdir(), key=lambda p: _image_set_sort_key(p.name))
        if d.is_dir() and is_image_set_dir(d.name)
    ]


def parse_new_listing_unit(
    product_code: str,
    unit_root: Path,
    image_set: str,
    excel_map: dict[str, SkuExcelInfo],
    defaults: dict[str, str],
    *,
    include_excel_missing_previews: bool = False,
) -> ListingUnit:
    label = f"{product_code}/{image_set}" if image_set else product_code
    family_code = ""
    try:
        if unit_root.parent.name == image_set and unit_root.parent.parent:
            family_code = unit_root.parent.parent.name
    except Exception:
        family_code = ""
    unit = ListingUnit(
        product_code=product_code,
        family_code=family_code,
        image_set=image_set,
        root=unit_root,
    )
    unit.carousel_images = _collect_numbered_images(unit_root, {1, 2, 3, 4, 5})
    unit.detail_images = _collect_numbered_images(unit_root, {7, 8, 9})

    if not unit.carousel_images:
        unit.warnings.append(f"[{label}] 未找到 1-5 主图")
    if not unit.detail_images:
        unit.warnings.append(f"[{label}] 未找到 7-9 详情图")

    for preview in sorted(
        _collect_flat_preview_images(unit_root),
        key=lambda p: _preview_sort_key(p, product_code),
    ):
        parsed = _parse_combo_code(preview.stem, product_code)
        if not parsed:
            unit.warnings.append(f"[{label}] 预览图文件名无法解析组合编码: {preview.name}")
            continue

        style_code, model_code = parsed
        style_name = lookup_sku_name(style_code, excel_map)
        model_name = lookup_sku_name(model_code, excel_map)
        if style_name is None:
            if not include_excel_missing_previews:
                unit.warnings.append(f"[{label}] Excel 无款式编码[{style_code}]，不上架 {preview.name}")
                continue
            unit.warnings.append(f"[{label}] Excel 无款式编码[{style_code}]，仅后台展示不上架 {preview.name}")
            style_name = style_code
        if model_name is None:
            if not include_excel_missing_previews:
                unit.warnings.append(f"[{label}] Excel 无型号编码[{model_code}]，不上架 {preview.name}")
                continue
            unit.warnings.append(f"[{label}] Excel 无型号编码[{model_code}]，仅后台展示不上架 {preview.name}")
            model_name = model_code

        if style_code not in unit.style_codes:
            unit.style_codes.append(style_code)
        if model_code not in unit.model_codes:
            unit.model_codes.append(model_code)

        inv, gp, sp = resolve_combo_pricing(style_code, model_code, excel_map, defaults)
        unit.skus.append(
            SkuCombo(
                style_code=style_code,
                style_name=style_name,
                model_code=model_code,
                model_name=model_name,
                combo_folder=preview.stem,
                preview_image=preview,
                inventory=inv,
                group_price=gp,
                single_price=sp,
            )
        )

    if not unit.skus:
        unit.warnings.append(f"[{label}] 未解析到任何有效 SKU")
    return unit


def parse_legacy_listing_unit(
    product_code: str,
    unit_root: Path,
    image_set: str,
    excel_map: dict[str, SkuExcelInfo],
    defaults: dict[str, str],
) -> ListingUnit:
    label = f"{product_code}/{image_set}" if image_set else product_code
    unit = ListingUnit(product_code=product_code, image_set=image_set, root=unit_root)
    unit.carousel_images = _collect_images(unit_root / "主图")
    unit.detail_images = _collect_images(unit_root / "详情页")

    style_dirs = sorted(
        (p for p in unit_root.iterdir() if p.is_dir() and p.name not in RESERVED_DIRS),
        key=lambda p: p.name,
    )

    for style_dir in style_dirs:
        if is_image_set_dir(style_dir.name):
            continue

        style_code = style_dir.name
        style_name = lookup_sku_name(style_code, excel_map)
        if style_name is None:
            unit.warnings.append(f"[{label}] Excel 无款式编码[{style_code}]，跳过")
            continue

        unit.style_codes.append(style_code)

        combo_dirs = sorted(
            (p for p in style_dir.iterdir() if p.is_dir()),
            key=lambda p: p.name,
        )
        for combo_dir in combo_dirs:
            parsed = _parse_combo_code(combo_dir.name, style_code)
            if not parsed:
                unit.warnings.append(f"[{label}] 组合文件夹格式错误: {combo_dir.name}")
                continue

            _, model_code = parsed
            model_name = lookup_sku_name(model_code, excel_map)
            if model_name is None:
                unit.warnings.append(f"[{label}] Excel 无型号编码[{model_code}]，跳过")
                continue

            preview = _find_preview_image(combo_dir)
            if not preview:
                unit.warnings.append(f"[{label}] 缺少预览图: {style_code}/{combo_dir.name}")
                continue

            if model_code not in unit.model_codes:
                unit.model_codes.append(model_code)

            inv, gp, sp = resolve_combo_pricing(style_code, model_code, excel_map, defaults)
            unit.skus.append(
                SkuCombo(
                    style_code=style_code,
                    style_name=style_name,
                    model_code=model_code,
                    model_name=model_name,
                    combo_folder=combo_dir.name,
                    preview_image=preview,
                    inventory=inv,
                    group_price=gp,
                    single_price=sp,
                )
            )

    if not unit.skus:
        unit.warnings.append(f"[{label}] 未解析到任何有效 SKU")
    return unit


def parse_listing_unit(
    product_code: str,
    unit_root: Path,
    image_set: str,
    excel_map: dict[str, SkuExcelInfo],
    defaults: dict[str, str],
    *,
    include_excel_missing_previews: bool = False,
) -> ListingUnit:
    if any(_numbered_image_index(p) is not None for p in unit_root.iterdir() if p.is_file()):
        return parse_new_listing_unit(
            product_code,
            unit_root,
            image_set,
            excel_map,
            defaults,
            include_excel_missing_previews=include_excel_missing_previews,
        )
    return parse_legacy_listing_unit(product_code, unit_root, image_set, excel_map, defaults)


def parse_all_listings(
    product_code: str,
    settings: YamlSettings,
    *,
    include_excel_missing_previews: bool = False,
) -> list[ListingUnit]:
    assets_root = settings.paths.resolve("assets_root")
    excel_map = _load_excel_map_with_fallback(settings, assets_root)
    defaults = settings.multi_spec.defaults

    new_units = _new_layout_units(product_code, assets_root)
    if new_units:
        return [
            parse_new_listing_unit(
                product_code,
                unit_root,
                image_set,
                excel_map,
                defaults,
                include_excel_missing_previews=include_excel_missing_previews,
            )
            for image_set, unit_root in new_units
        ]

    product_root = assets_root / product_code
    if not product_root.is_dir():
        return []

    image_sets = list_image_sets(product_code, assets_root)
    if image_sets:
        return [
            parse_listing_unit(
                product_code,
                product_root / image_set,
                image_set,
                excel_map,
                defaults,
                include_excel_missing_previews=include_excel_missing_previews,
            )
            for image_set in image_sets
        ]
    return [
        parse_listing_unit(
            product_code,
            product_root,
            "",
            excel_map,
            defaults,
            include_excel_missing_previews=include_excel_missing_previews,
        )
    ]


def parse_listing_target(
    target: str,
    settings: YamlSettings,
    *,
    include_excel_missing_previews: bool = False,
) -> list[ListingUnit]:
    selector = _parse_listing_target_selector(target)
    if not selector:
        return parse_all_listings(
            target,
            settings,
            include_excel_missing_previews=include_excel_missing_previews,
        )

    family_code, image_set = selector
    assets_root = settings.paths.resolve("assets_root")
    excel_map = _load_excel_map_with_fallback(settings, assets_root)
    defaults = settings.multi_spec.defaults
    set_dir = assets_root / family_code / image_set
    if not set_dir.is_dir():
        return []

    units: list[ListingUnit] = []
    for code_dir in sorted((p for p in set_dir.iterdir() if p.is_dir()), key=lambda p: p.name):
        if not _code_dir_has_numbered_images(code_dir):
            continue
        units.append(
            parse_new_listing_unit(
                code_dir.name,
                code_dir,
                image_set,
                excel_map,
                defaults,
                include_excel_missing_previews=include_excel_missing_previews,
            )
        )
    return units


def list_listing_targets(settings: YamlSettings) -> list[ListingTargetOption]:
    assets_root = settings.paths.resolve("assets_root")
    return _new_layout_listing_targets(assets_root)


def list_product_codes(settings: YamlSettings) -> list[str]:
    assets_root = settings.paths.resolve("assets_root")
    new_codes = _new_layout_product_codes(assets_root)
    if new_codes:
        return new_codes
    if not assets_root.is_dir():
        return []
    return sorted(d.name for d in assets_root.iterdir() if d.is_dir() and not d.name.startswith("."))


class CatalogSkill(Skill):
    name = "parse_catalog"
    description = "解析多规格素材目录，返回可上架单元列表"

    def as_tool(self) -> StructuredTool:
        def _run(product_code: str) -> str:
            from pdd_agent.core.config import load_yaml_settings

            settings = load_yaml_settings()
            units = parse_all_listings(product_code, settings)
            lines = [f"{u.listing_label}: {len(u.skus)} skus" for u in units]
            return "\n".join(lines) if lines else "no units"

        return StructuredTool.from_function(
            name=self.name,
            description=self.description,
            func=_run,
        )
